"""Review system L1-L3 (A6): record formats, hash binding, sampling rules and the published queue."""
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'tools'))
import review_system as rs  # noqa: E402
import review as tool  # noqa: E402
import select_prototype as sp  # noqa: E402
import review_sheet as sheet  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

QUEUE = json.loads((ROOT / 'docs/data/review-queue.json').read_text(encoding='utf-8'))
CATALOG = json.loads((ROOT / 'docs/data/catalog.json').read_text(encoding='utf-8'))
QUESTIONS = rs.load_questions()
SELECTION = json.loads((ROOT / 'docs/data/prototype-selection.json').read_text(encoding='utf-8'))


def rec(**k):
    base = dict(level='L1', unit='named_bond_angle', decision='approved', date='2026-09-30', reviewer=dict(kind='human', name='Reviewer'),
                bindings={'registry_row': 'a', 'task_families/local_geometry.py': 'b'})
    base.update(k)
    return base


def l3(instance, verdicts):
    return dict(instance=instance, reviewer=dict(kind='model', name='some-model'), questions_sha256=QUESTIONS['_sha256'],
                answers={q['id']: dict(verdict=v, reason='r') for q, v in zip(QUESTIONS['questions'], verdicts)})


class RecordTests(unittest.TestCase):
    def test_l1_l2_need_a_named_human(self):
        rs.validate_record(rec())
        for bad in (rec(reviewer=dict(kind='model', name='x')), rec(reviewer=dict(kind='human', name=' ')), rec(decision='ok'),
                    rec(level='L3'), rec(bindings={}), rec(date='30/09/2026')):
            with self.assertRaises(ValueError):
                rs.validate_record(bad)

    def test_changed_code_makes_a_record_stale(self):
        records = dict(L1=[rec()], L2=[], L3=[])
        self.assertEqual(rs.unit_state('named_bond_angle', 'L1', rec()['bindings'], records)[0], 'approved')
        changed = dict(rec()['bindings'], **{'task_families/local_geometry.py': 'c'})
        self.assertEqual(rs.unit_state('named_bond_angle', 'L1', changed, records)[0], 'stale')
        self.assertEqual(rs.unit_state('other', 'L1', changed, records)[0], 'pending')
        later = rec(decision='revise', date='2026-10-01')
        self.assertEqual(rs.unit_state('named_bond_angle', 'L1', rec()['bindings'], dict(L1=[rec(), later], L2=[], L3=[]))[0], 'revise')

    def test_l3_results_must_answer_the_fixed_four_questions(self):
        rs.validate_l3(l3('X1', ['pass'] * 4), QUESTIONS)
        bad = l3('X1', ['pass'] * 4)
        bad['answers'].pop(QUESTIONS['questions'][0]['id'])
        with self.assertRaises(ValueError):
            rs.validate_l3(bad, QUESTIONS)
        with self.assertRaises(ValueError):
            rs.validate_l3(dict(l3('X1', ['pass'] * 4), questions_sha256='0' * 64), QUESTIONS)
        with self.assertRaises(ValueError):
            rs.validate_l3(l3('X1', ['pass', 'maybe', 'pass', 'pass']), QUESTIONS)

    def test_records_are_append_only_and_need_a_reviewer(self):
        with tempfile.TemporaryDirectory() as tmp:
            old = rs.REVIEWS
            rs.REVIEWS = Path(tmp)
            try:
                p1 = tool.record('L1', 'named_bond_angle', 'Reviewer', 'approved', 'ok', '2026-09-30')
                p2 = tool.record('L1', 'named_bond_angle', 'Reviewer', 'revise', 'fix', '2026-09-30')
                self.assertNotEqual(p1, p2)
                self.assertEqual(json.loads(p1.read_text(encoding='utf-8'))['decision'], 'approved')
                with self.assertRaises(SystemExit):
                    tool.record('L1', 'named_bond_angle', ' ', 'approved', '', '2026-09-30')
                with self.assertRaises(SystemExit):
                    tool.record('L1', 'no_such_family', 'Reviewer', 'approved', '', '2026-09-30')
                loaded = rs.load_records(Path(tmp))
                self.assertEqual(len(loaded['L1']), 2)
            finally:
                rs.REVIEWS = old


class SamplingTests(unittest.TestCase):
    ITEMS = [dict(id='A%02d' % i, unit='u1', source='s1', cell='chemistry 1-10', l0=True) for i in range(30)] + \
            [dict(id='B%02d' % i, unit='u2', source='s2', cell='quantum 0.1-1', l0=True) for i in range(5)]

    def test_first_batch_is_checked_in_full(self):
        plan = rs.plan_l3(self.ITEMS, [], 'seed')
        self.assertEqual(sum(p['sample'] for p in plan), 35)

    def test_stable_strata_get_max_two_or_ten_percent(self):
        history = [dict(results=[dict(unit='u1', source='s1', final='pass'), dict(unit='u2', source='s2', final='pass')])]
        plan = {p['unit']: p for p in rs.plan_l3(self.ITEMS, history, 'seed')}
        self.assertEqual(plan['u1']['sample'], 3)            # ceil(10% of 30)
        self.assertEqual(plan['u2']['sample'], 2)            # minimum 2
        self.assertEqual(plan['u1']['ids'], rs.plan_l3(self.ITEMS, history, 'seed')[0]['ids'])   # deterministic
        self.assertNotEqual(plan['u1']['ids'], {p['unit']: p for p in rs.plan_l3(self.ITEMS, history, 'other')}['u1']['ids'])

    def test_flags_and_calibration_go_to_a_human(self):
        results = [l3('P%02d' % i, ['pass'] * 4) for i in range(20)] + [l3('F1', ['pass', 'flag', 'pass', 'pass'])]
        tri = rs.triage(results, 'seed')
        self.assertEqual(tri['flagged'], ['F1'])
        self.assertEqual(len(tri['calibration']), 3)          # max(3, 10% of 20 passes)
        self.assertTrue(all(i.startswith('P') for i in tri['calibration']))

    def test_status_flow(self):
        item = dict(id='X', unit='u', l0=True)
        self.assertEqual(rs.instance_stage(item, {}, {}), 'L0')
        self.assertEqual(rs.instance_stage(item, {'u': ('approved', 'pending')}, {}), 'L0')
        self.assertEqual(rs.instance_stage(item, {'u': ('approved', 'approved')}, {}), 'L1+L2')
        self.assertEqual(rs.instance_stage(item, {'u': ('approved', 'approved')}, {'X': 'pass'}), 'formal')
        self.assertEqual(rs.instance_stage(item, {'u': ('approved', 'stale')}, {'X': 'pass'}), 'L0')
        self.assertEqual(rs.instance_stage(dict(item, l0=False), {}, {}), 'not-verified')


class SelectionTests(unittest.TestCase):
    def test_published_selection_is_current_and_deterministic(self):
        self.assertEqual(sp.build(), SELECTION)               # rerun tools/select_prototype.py after catalog changes
        self.assertEqual(SELECTION['status'], 'proposal')

    def test_ten_per_cell_and_everything_else_is_reserve(self):
        active = sorted(q['id'] for q in CATALOG['questions'] if q['lifecycle'] == 'active')
        self.assertEqual(len(SELECTION['cells']), 16)
        for c in SELECTION['cells']:
            self.assertEqual(len(c['selected']), 10, c['cell'])
            self.assertTrue(all(s['form'] != 'legacy' or s['unit'].startswith('legacy:') for s in c['selected']))
        both = sorted([s['id'] for c in SELECTION['cells'] for s in c['selected']] + [r['id'] for c in SELECTION['cells'] for r in c['reserve']])
        self.assertEqual(both, active)
        self.assertEqual(sorted(SELECTION['selected_ids']), sorted(s['id'] for c in SELECTION['cells'] for s in c['selected']))

    def test_checked_questions_come_first(self):
        for c in SELECTION['cells']:
            legacy_used = any(s['form'] == 'legacy' for s in c['selected'])
            checked_left = [r['id'] for r in c['reserve'] if r['form'] != 'legacy' and r['reason'].startswith('本格')]
            self.assertFalse(legacy_used and checked_left, c['cell'])

    def test_abilities_are_balanced_within_the_tier(self):
        q = lambda i, a, form='family', unit='u', pending=None: dict(id=i, ability=a, form=form, unit=unit, instance=i, source='s',
                                                                   pending=pending, tier=0 if form != 'legacy' else 1, l0=True)
        rule = sp.read('templates/prototype-selection-rule.json')
        pool = [q('P%d' % i, 'perception') for i in range(8)] + [q('I%d' % i, 'inference') for i in range(8)] + [q('D%d' % i, 'design') for i in range(2)]
        got = sp.pick(pool, rule, [])
        self.assertEqual([sum(c['ability'] == a for c in got) for a in rule['abilities']], [4, 4, 2])
        pool = [q('D1', 'design', pending='x'), q('D2', 'design'), q('P1', 'perception')]
        self.assertEqual([c['id'] for c in sp.pick(pool, dict(rule, per_cell=2), [])], ['P1', 'D2'])   # pending groups last

    def test_legacy_record_with_checked_version_is_reviewed_with_its_family(self):
        q = dict(id='X', family='global_extent', familyVersions=[dict(instance='FC-X', family='extent_choice_v2', independentCheck='pass')])
        self.assertEqual(rs.review_form(q), dict(form='family-version', unit='extent_choice_v2', instance='FC-X', l0=True))
        q['familyVersions'][0]['independentCheck'] = 'fail'
        self.assertEqual(rs.review_form(q)['unit'], 'legacy:global_extent')


class PublishedQueueTests(unittest.TestCase):
    def test_queue_covers_every_selected_record_once(self):
        planned = sorted(i for p in QUEUE['l3_plan'] for i in p['ids'])
        self.assertEqual(QUEUE['summary']['scope'], 'selection')
        self.assertEqual(planned, sorted(SELECTION['selected_ids']))      # first batch: everything, each once
        self.assertEqual(sorted(QUEUE['stages']), sorted(SELECTION['selected_ids']))
        self.assertEqual(QUEUE['summary']['model_calls'], 0)
        self.assertEqual(QUEUE['questions_sha256'], QUESTIONS['_sha256'])
        self.assertEqual(QUEUE['summary']['units'], SELECTION['summary']['review_units'])

    def test_every_unit_has_a_card_and_bindings(self):
        units = {u['unit']: u for u in QUEUE['units']}
        for q in CATALOG['questions']:
            if q['id'] in SELECTION['selected_ids']:
                self.assertIn(rs.review_form(q)['unit'], units)
        for u in QUEUE['units']:
            self.assertTrue(u['bindings']['L1'].get('registry_row'), u['unit'])
            if u['kind'] == 'family':
                self.assertTrue(u['checker'] and all(u['bindings']['L2'].values()), u['unit'])
                if u['instances']:
                    self.assertGreaterEqual(len(u['card']['samples']), 1, u['unit'])

    def test_no_review_is_recorded_on_anyones_behalf(self):
        self.assertFalse(any((ROOT / 'reviews' / lvl).exists() and any((ROOT / 'reviews' / lvl).iterdir()) for lvl in ('L1', 'L2', 'L3')))
        self.assertEqual(QUEUE['summary']['stages']['formal'], 0)

    def test_prompt_contains_packets_and_the_four_questions(self):
        text = rs.l3_prompt('X1', 'QUESTION TEXT', dict(correct_label='B'), QUESTIONS)
        for q in QUESTIONS['questions']:
            self.assertIn(q['id'], text)
        self.assertIn('=== STUDENT PACKET ===\nQUESTION TEXT', text)
        self.assertIn('"correct_label": "B"', text)


class SheetTests(unittest.TestCase):
    def fill(self, path, edits):
        wb = load_workbook(path)
        for (title, row), values in edits.items():
            ws = wb[title]
            head = [c.value for c in ws[1]]
            for k, v in values.items():
                ws.cell(row=row, column=head.index(k) + 1, value=v)
        wb.save(path)

    def test_export_has_one_row_per_unit_with_bindings(self):
        with tempfile.TemporaryDirectory() as tmp:
            out, n = sheet.export(Path(tmp) / 's.xlsx')
            self.assertEqual(n, len(QUEUE['units']))
            wb = load_workbook(out)
            for title, level in (('L1 题型审定', 'L1'), ('L2 检查器审阅', 'L2')):
                ws = wb[title]
                head = [c.value for c in ws[1]]
                got = [json.loads(r[head.index('绑定（勿改）')]) for r in ws.iter_rows(min_row=2, values_only=True)]
                self.assertEqual([g['unit'] for g in got], [u['unit'] for u in QUEUE['units']])
                self.assertEqual([g['bindings'] for g in got], [u['bindings'][level] for u in QUEUE['units']])

    def test_import_validates_everything_before_writing(self):
        with tempfile.TemporaryDirectory() as tmp:
            old = rs.REVIEWS
            rs.REVIEWS = Path(tmp) / 'reviews'
            try:
                path = Path(tmp) / 's.xlsx'
                sheet.export(path)
                ok = {'决定（通过 / 需修改）': '通过', '审核人姓名': 'Reviewer', '日期（YYYY-MM-DD）': '2026-10-03', '意见 / 修改要求': 'fine'}
                self.fill(path, {('L1 题型审定', 2): ok, ('L2 检查器审阅', 3): dict(ok, **{'决定（通过 / 需修改）': '需修改'})})
                self.assertEqual(sheet.do_import(path, dry_run=True)['to_write'], 2)
                self.assertFalse(rs.REVIEWS.exists())                         # dry run writes nothing
                bad = Path(tmp) / 'bad.xlsx'
                bad.write_bytes(path.read_bytes())
                self.fill(bad, {('L1 题型审定', 4): dict(ok, **{'审核人姓名': ''}),
                                ('L1 题型审定', 5): dict(ok, **{'绑定（勿改）': json.dumps(dict(level='L1', unit=QUEUE['units'][3]['unit'], bindings={'registry_row': 'old'}))})})
                with self.assertRaises(SystemExit) as e:
                    sheet.do_import(bad)
                self.assertIn('缺审核人姓名', str(e.exception))
                self.assertIn('绑定不一致', str(e.exception))
                self.assertFalse(rs.REVIEWS.exists())                         # one bad row: nothing written
                res = sheet.do_import(path)
                self.assertEqual(len(res['written']), 2)
                recs = rs.load_records(rs.REVIEWS)
                self.assertEqual(recs['L1'][0]['unit'], QUEUE['units'][0]['unit'])
                self.assertEqual(recs['L2'][0]['decision'], 'revise')
                self.assertEqual(recs['L1'][0]['reviewer'], dict(kind='human', name='Reviewer'))
                again = sheet.do_import(path)                                 # same sheet twice: no duplicates
                self.assertEqual((again['to_write'], again['already_imported']), (0, 2))
            finally:
                rs.REVIEWS = old


if __name__ == '__main__':
    unittest.main()
