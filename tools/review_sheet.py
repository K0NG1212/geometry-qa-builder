"""L1 / L2 review sheet (Excel) for human reviewers, and import of the filled sheet into reviews/.

  python tools/review_sheet.py export [--all] [--out docs/data/review-sheet.xlsx]
  python tools/review_sheet.py import --file filled.xlsx [--dry-run]

Export: one row per review unit on the L1 and L2 sheets (default scope: the 160 selection), with the family card
summary, two samples, links to the review page / workbench / code, a decision dropdown, reviewer, date and notes.
Each row carries the binding hashes current at export time.
Import: validates every filled row first (decision, named reviewer, date, unit exists, bindings unchanged since
export) and writes nothing if any row fails; then writes the records with tools/review.py (append-only). Rows that
were already imported are skipped. A reviewer name must be the person who made the decision, never an AI assistant.
"""
import argparse
import datetime
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'tools'))
import review_system as rs  # noqa: E402
import review as tool  # noqa: E402

SITE = 'https://k0ng1212.github.io/geometry-qa-builder/'
REPO = 'https://github.com/K0NG1212/geometry-qa-builder/blob/main/'
CHOICES = {'通过': 'approved', '需修改': 'revise', 'approved': 'approved', 'revise': 'revise'}
ABILITY = {'perception': '感知', 'inference': '推断', 'design': '设计'}
STATE = {'pending': '待审', 'approved': '通过', 'revise': '需修改', 'stale': '已失效'}
INPUT = ['决定（通过 / 需修改）', '审核人姓名', '日期（YYYY-MM-DD）', '意见 / 修改要求']
GREY = PatternFill('solid', fgColor='EEEEEE')
YELLOW = PatternFill('solid', fgColor='FFF7D6')
CARD_KEYS = (('name', '名称'), ('concept', '考点'), ('physical_conditions', '物理条件'), ('reasoning_scale', '推理尺度'),
             ('answer_method', '答案方法'), ('distractor_mechanisms', '干扰机制'), ('limits', '限制'))


def text(v):
    if v is None:
        return ''
    if isinstance(v, (list, tuple)):
        return '；'.join(text(x) for x in v)
    if isinstance(v, dict):
        return '；'.join('%s: %s' % (k, text(x)) for k, x in v.items())
    return str(v)


def card_text(card):
    t = card['template']
    return '\n'.join('【%s】%s' % (label, text(t[k])) for k, label in CARD_KEYS if t.get(k) not in (None, '', []))


def sample_text(s):
    if not s:
        return ''
    if s.get('instance'):
        opts = '\n'.join('%s. %s %s' % (o['label'], o['value'], o.get('unit') or '') for o in s['options'])
        audit = '\n'.join('%s · %s：%s' % (a['label'], a['rule'] or '', a['reason'] or '') for a in s['audit'])
        return '%s（正确 %s）\n%s\n\n%s\n\n附件：%s\n\n干扰项：\n%s' % (s['instance'], s['correct_label'], s['question'], opts, '、'.join(s['inputs']), audit)
    return '%s\n%s\n\n答案：%s\n\n验证：%s' % (s['record'], s['question'] or '', text(s.get('answer')), text(s.get('validation')))


def sample_link(s):
    if not s:
        return ''
    if s.get('instance'):
        return SITE + 'templates.html?instance=%s#workbench' % s['instance']
    return SITE + 'index.html?qa=%s#questions' % s['record']


COLUMNS = {
    'L1': ['单元', '类型', '能力', '选中题数', '注册表', '题型卡片', '样例 1', '样例 1 链接', '样例 2', '样例 2 链接', '审核页', '生成代码', '当前状态'],
    'L2': ['单元', '类型', '能力', '选中题数', '检查器', '检查器代码', '生成代码', '检查方法（注册表）', '审核页', '当前状态'],
}
WIDTH = {'题型卡片': 60, '样例 1': 60, '样例 2': 60, '检查方法（注册表）': 50, '意见 / 修改要求': 40, '单元': 26}


def rows(level, queue):
    for u in queue['units']:
        page = SITE + 'review.html?unit=%s' % u['unit']
        common = [u['unit'], '题型族' if u['kind'] == 'family' else '旧模板', ABILITY.get(u['ability'], u['ability']), u['instances']]
        gen = REPO + u['generator'] if u.get('generator') else ''
        if level == 'L1':
            s = u['card']['samples'] + [None, None]
            yield u, common + [u.get('registry_id') or '', card_text(u['card']), sample_text(s[0]), sample_link(s[0]), sample_text(s[1]), sample_link(s[1]),
                               page, gen, STATE.get(u['L1'], u['L1'])]
        else:
            if u['kind'] == 'family':
                checker, code = '%s:%s' % (u['checker'], u['checker_function']), REPO + u['checker']
            else:
                checker, code = '（旧模板：审阅计算代码）', '\n'.join(REPO + c for c in u.get('code', []))
            t = u['card']['template']
            method = '\n'.join('【%s】%s' % (k, text(t[k])) for k in ('validation', 'four_choice_validation', 'answer_method') if t.get(k))
            yield u, common + [checker, code, gen, method, page, STATE.get(u['L2'], u['L2'])]


def export(out):
    """Write the workbook from the published queue (docs/data/review-queue.json; the CLI refreshes it first)."""
    queue = json.loads((ROOT / 'docs/data/review-queue.json').read_text(encoding='utf-8'))
    scope = queue['summary']['scope']
    wb = Workbook()
    info = wb.active
    info.title = '说明'
    lines = ['GeoBench 题型审核表（L1 题型审定 / L2 检查器审阅）',
             '导出日期：%s；范围：%s；审核单元 %d 个。' % (queue['date'], '160 道原型挑选（提议，待教授确认）' if scope == 'selection' else '全部 active 题', len(queue['units'])),
             '',
             '填写：只填黄色四列——决定（下拉选“通过”或“需修改”）、审核人姓名、日期、意见。灰色列请勿修改；最右侧“绑定”列用于核对代码是否在导出后改动。',
             'L1：读题型卡片和两个样例（可点链接看完整题目、附件与工作台），判断考点、题干条件、物理假设、干扰项机制与限制是否成立。',
             'L2：读独立检查器代码（链接到 GitHub），判断它是否用与生成器不同的方法、从题目本身读取参数并重算答案。',
             '“需修改”请在意见里写明要改什么；修改后该题型整体重新生成，不逐题手改。',
             '审核人必须是作出判断的人本人；AI 助手不能代填或代为批准。',
             '没把握的行可以留空，导入时只处理填了决定的行。',
             '',
             '交回后导入：python tools/review_sheet.py import --file <文件名>.xlsx （先加 --dry-run 只检查不写入）',
             '导入会逐行检查：决定合法、审核人非空、日期格式、单元存在、绑定哈希与当前代码一致（代码已改则拒收该行，需要重新导出）。任一行不合格则整份不写入。',
             '',
             '网站：' + SITE + 'review.html']
    for i, line in enumerate(lines, 1):
        info.cell(row=i, column=1, value=line)
    info['A1'].font = Font(bold=True, size=14)
    info.column_dimensions['A'].width = 120
    for level, title in (('L1', 'L1 题型审定'), ('L2', 'L2 检查器审阅')):
        ws = wb.create_sheet(title)
        head = COLUMNS[level] + INPUT + ['绑定（勿改）']
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
        dv = DataValidation(type='list', formula1='"通过,需修改"', allow_blank=True)
        ws.add_data_validation(dv)
        for u, values in rows(level, queue):
            ws.append(values + ['', '', '', '', json.dumps(dict(level=level, unit=u['unit'], bindings=u['bindings'][level]), ensure_ascii=False, sort_keys=True)])
            r = ws.max_row
            dv.add(ws.cell(row=r, column=len(COLUMNS[level]) + 1))
            for col in range(1, len(head) + 1):
                cell = ws.cell(row=r, column=col)
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                cell.fill = YELLOW if len(COLUMNS[level]) < col <= len(COLUMNS[level]) + len(INPUT) else GREY
                if isinstance(cell.value, str) and cell.value.startswith('https://') and '\n' not in cell.value:
                    cell.hyperlink = cell.value
                    cell.font = Font(color='0563C1', underline='single')
        for i, h in enumerate(head, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = WIDTH.get(h, 14 if i <= len(COLUMNS[level]) else 18)
        ws.freeze_panes = 'B2'
    sel = wb.create_sheet('160 道挑选')
    path = ROOT / 'docs/data/prototype-selection.json'
    if path.exists():
        data = json.loads(path.read_text(encoding='utf-8'))
        sel.append(['题号', '格', '能力', '形式', '审核单元', '题型族实例', '待教授决定', '链接'])
        for c in data['cells']:
            for x in c['selected']:
                sel.append([x['id'], c['cell'], ABILITY[x['ability']], {'family': '题型族实例', 'family-version': '旧题的四选一版本', 'legacy': '旧模板题'}[x['form']],
                            x['unit'], x['instance'] or '', data['pending_labels'].get(x['pending'], '') if x['pending'] else '', SITE + 'index.html?qa=%s#questions' % x['id']])
        for col, w in zip('ABCDEFGH', (10, 18, 8, 16, 28, 26, 40, 60)):
            sel.column_dimensions[col].width = w
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out, len(queue['units'])


def parse_date(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    v = str(v or '').strip()
    return v if re.fullmatch(r'\d{4}-\d{2}-\d{2}', v) else None


def filled_rows(path):
    wb = load_workbook(path)
    for title, level in (('L1 题型审定', 'L1'), ('L2 检查器审阅', 'L2')):
        if title not in wb.sheetnames:
            continue
        ws = wb[title]
        head = [c.value for c in ws[1]]
        idx = {h: i for i, h in enumerate(head)}
        for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            decision = str(row[idx[INPUT[0]]] or '').strip()
            if not decision:
                continue
            yield dict(sheet=title, row=n, level=level, decision=decision, reviewer=str(row[idx[INPUT[1]]] or '').strip(),
                       date=row[idx[INPUT[2]]], notes=str(row[idx[INPUT[3]]] or '').strip(), binding=row[idx['绑定（勿改）']])


def check(path):
    """Validate every filled row; returns (records to write, errors, already imported)."""
    registry, catalog, workbench, units = tool.context()
    existing = rs.load_records()
    todo, errors, done = [], [], []
    for r in filled_rows(path):
        where = '%s 第 %d 行' % (r['sheet'], r['row'])
        try:
            b = json.loads(r['binding'] or '')
        except ValueError:
            errors.append(where + '：绑定列缺失或被改动')
            continue
        unit = b.get('unit')
        if b.get('level') != r['level'] or unit not in units:
            errors.append(where + '：单元或层级不符（%s）' % unit)
            continue
        if r['decision'] not in CHOICES:
            errors.append(where + '：决定须为“通过”或“需修改”，现为 %r' % r['decision'])
        if not r['reviewer']:
            errors.append(where + '：缺审核人姓名')
        date = parse_date(r['date'])
        if not date:
            errors.append(where + '：日期须为 YYYY-MM-DD')
        current = rs.bindings(units[unit], registry)[r['level']]
        if b.get('bindings') != current:
            errors.append(where + '：%s 的代码或注册表在导出后已改动（绑定不一致），请重新导出后再审' % unit)
        if errors and errors[-1].startswith(where):
            continue
        rec = dict(level=r['level'], unit=unit, decision=CHOICES[r['decision']], reviewer=r['reviewer'], date=date, notes=r['notes'])
        same = [x for x in existing[r['level']] if x['unit'] == unit and x['decision'] == rec['decision'] and x['reviewer']['name'] == rec['reviewer']
                and x['date'] == date and x.get('notes', '') == rec['notes'] and x['bindings'] == current]
        (done if same else todo).append(rec)
    return todo, errors, done


def do_import(path, dry_run=False):
    todo, errors, done = check(path)
    if errors:
        raise SystemExit('未写入任何记录：\n' + '\n'.join(errors))
    written = []
    if not dry_run:
        for rec in todo:
            written.append(tool.record(rec['level'], rec['unit'], rec['reviewer'], rec['decision'], rec['notes'], rec['date']))
    return dict(to_write=len(todo), written=[str(p) for p in written], already_imported=len(done), dry_run=dry_run)


if __name__ == '__main__':
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest='cmd', required=True)
    e = sub.add_parser('export')
    e.add_argument('--out', default=str(ROOT / 'docs/data/review-sheet.xlsx'))
    e.add_argument('--all', action='store_true')
    i = sub.add_parser('import')
    i.add_argument('--file', required=True)
    i.add_argument('--dry-run', action='store_true')
    a = p.parse_args()
    if a.cmd == 'export':
        tool.queue(datetime.date.today().isoformat(), 'all' if a.all else 'selection')
        out, n = export(a.out)
        print(out, n, 'units')
    else:
        print(json.dumps(do_import(a.file, a.dry_run), ensure_ascii=False, indent=1))
